#!/usr/bin/env python3
"""docx 中文学术后处理：eastAsia 字体强制、标题自动编号、目录(TOC)域、图表 SEQ 题注。

开源 Office MCP（python-docx 系）共同的盲区：只写 latin 字体槽位（中文回落默认字体）、
无 TOC 工具、无题注编号。本脚本用 Word/WPS 原生支持的标准域补齐这三项。

用法:
    python postprocess_docx.py 报告.docx                        # 全默认：字体+编号+目录+题注
    python postprocess_docx.py 报告.docx --no-toc --number-levels 3
    python postprocess_docx.py 报告.docx --heading-font 黑体 --body-font 宋体

约定（由 SKILL.md 规定上游写法）:
    - 图注:  pandoc 图片 `![图注](路径)` 生成的 Image Caption 段落，或紧贴图片的 "图：xxx" 段落
    - 表注:  表格正上方一行 "表：xxx"
幂等: 重复运行不会二次编号/二次插目录。

输出 JSON 摘要；就地修改文件。退出码: 成功 0。
"""
from __future__ import annotations

import argparse
import json
import re
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

CJK_RE = re.compile(r"[\u4e00-\u9fff]")
NUM_PREFIX_RE = re.compile(r"^\s*(\d+(?:\.\d+)*)[\s.、]+")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
CAPTION_SIZE = None  # 题注字号沿用样式，不在此强设


def set_rfonts(rpr, ea: str | None, latin: str | None) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    if latin:
        rfonts.set(qn("w:ascii"), latin)
        rfonts.set(qn("w:hAnsi"), latin)
    if ea:
        rfonts.set(qn("w:eastAsia"), ea)


def style_is_heading(name: str) -> bool:
    return name.startswith("Heading") or name in ("Title", "TOC Heading")


# ---------------------------------------------------------------- 字体
def apply_fonts(doc: Document, heading_font: str, body_font: str, latin: str) -> dict:
    stats = {"styles_touched": 0, "runs_cjk": 0, "runs_cjk_fixed": 0}

    # 1) 文档默认字体（docDefaults）：决定没有显式字体时的回落
    styles_el = doc.styles.element
    dd_rpr = styles_el.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    if dd_rpr is None:
        dd = styles_el.find(qn("w:docDefaults"))
        if dd is not None:
            dd_rpr = OxmlElement("w:rPr")
            dd.append(dd_rpr)
    if dd_rpr is not None:
        set_rfonts(dd_rpr, body_font, latin)

    # 2) 样式层：标题类样式给标题字体，其余给正文字体
    for style in doc.styles:
        if style.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            continue
        ea = heading_font if style_is_heading(style.name) else body_font
        try:
            style.font.name = latin  # 写 ascii/hAnsi，同时确保 rPr 存在
        except Exception:
            continue
        rpr = style.element.find(qn("w:rPr"))
        if rpr is not None:
            set_rfonts(rpr, ea, latin)
            stats["styles_touched"] += 1

    # 3) run 层兜底：所有含中文的 run 显式写 eastAsia（优先级最高，保证显示）
    def fix_runs(paragraphs):
        for p in paragraphs:
            heading = False
            try:
                heading = style_is_heading(p.style.name) if p.style is not None else False
            except Exception:
                pass
            for run in p.runs:
                if not run.text or not CJK_RE.search(run.text):
                    continue
                stats["runs_cjk"] += 1
                rpr = run._r.get_or_add_rPr()
                set_rfonts(rpr, heading_font if heading else body_font, latin)
                stats["runs_cjk_fixed"] += 1

    fix_runs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                fix_runs(cell.paragraphs)
    return stats


# ---------------------------------------------------------------- 标题编号
def number_headings(doc: Document, max_level: int) -> dict:
    counters = {1: 0, 2: 0, 3: 0}
    numbered = 0
    for p in doc.paragraphs:
        try:
            name = p.style.name if p.style is not None else ""
        except Exception:
            continue
        m = re.match(r"^Heading ([1-3])$", name)
        if not m:
            continue
        level = int(m.group(1))
        if level > max_level:
            continue
        text = p.text.strip()
        if not text:
            continue
        existing = NUM_PREFIX_RE.match(text)
        if existing:
            parts = [int(x) for x in existing.group(1).split(".")]
            for i, v in enumerate(parts, start=1):
                counters[i] = v
            continue
        counters[level] += 1
        for lower in range(level + 1, 4):
            counters[lower] = 0
        num = ".".join(str(counters[i]) for i in range(1, level + 1))
        target = next((r for r in p.runs if r.text), None)
        if target is None:
            continue
        target.text = f"{num} {target.text}"
        numbered += 1
    return {"headings_numbered": numbered}


# ---------------------------------------------------------------- 目录域
def _fld_char_run(p: Paragraph, kind: str) -> None:
    run = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), kind)
    if kind == "begin":
        fld.set(qn("w:dirty"), "true")
    run._r.append(fld)


def _instr_run(p: Paragraph, code: str) -> None:
    run = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(XML_SPACE, "preserve")
    instr.text = code
    run._r.append(instr)


def insert_toc(doc: Document, toc_levels: str) -> dict:
    body_xml = doc.element.body.xml
    if "TOC" in body_xml and "instrText" in body_xml:
        return {"toc_inserted": False, "toc_note": "已存在目录域，跳过"}
    first_heading = next(
        (p for p in doc.paragraphs
         if p.style is not None and p.style.name.startswith("Heading")), None)
    if first_heading is None:
        return {"toc_inserted": False, "toc_note": "未找到 Heading 样式段落，跳过"}

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("目  录")
    run.bold = True
    first_heading._p.addprevious(title._p)

    field = doc.add_paragraph()
    _fld_char_run(field, "begin")
    _instr_run(field, f' TOC \\o "{toc_levels}" \\h \\z \\u ')
    _fld_char_run(field, "separate")
    field.add_run("（目录将在打开文档时自动更新）")
    _fld_char_run(field, "end")
    first_heading._p.addprevious(field._p)

    pagebreak = doc.add_paragraph()
    br_run = pagebreak.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    br_run._r.append(br)
    first_heading._p.addprevious(pagebreak._p)

    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)
    return {"toc_inserted": True, "toc_levels": toc_levels}


# ---------------------------------------------------------------- 图表题注
def _clear_paragraph(p: Paragraph) -> None:
    for child in list(p._p):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            p._p.remove(child)


def _add_seq_field(p: Paragraph, seq_name: str) -> None:
    _fld_char_run(p, "begin")
    _instr_run(p, f" SEQ {seq_name} \\* ARABIC ")
    _fld_char_run(p, "separate")
    p.add_run("1")
    _fld_char_run(p, "end")


def _has_seq(p: Paragraph) -> bool:
    return "instrText" in p._p.xml and "SEQ" in p._p.xml


def _format_caption(p: Paragraph) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.bold = True


def _rewrite_caption(p: Paragraph, prefix: str, label: str) -> None:
    _clear_paragraph(p)
    p.add_run(f"{prefix} ")
    _add_seq_field(p, prefix)
    p.add_run(f"  {label}")
    _format_caption(p)


def number_captions(doc: Document) -> dict:
    stats = {"fig_captions": 0, "table_captions": 0}
    body = doc.element.body
    children = list(body.iterchildren())
    para_cache: dict[int, Paragraph] = {}

    def as_para(el) -> Paragraph | None:
        if el.tag != qn("w:p"):
            return None
        key = id(el)
        if key not in para_cache:
            para_cache[key] = Paragraph(el, doc)
        return para_cache[key]

    has_drawing = lambda el: el is not None and el.find(".//" + qn("w:drawing")) is not None

    # 先处理表注："表：xxx" 段落且后随表格
    for i, el in enumerate(children):
        p = as_para(el)
        if p is None or _has_seq(p):
            continue
        text = p.text.strip()
        if not re.match(r"^表[：:]", text):
            continue
        nxt = children[i + 1] if i + 1 < len(children) else None
        if nxt is not None and nxt.tag == qn("w:tbl"):
            _rewrite_caption(p, "表", re.sub(r"^表[：:]\s*", "", text))
            p.paragraph_format.keep_with_next = True  # 题注与表格不分离
            stats["table_captions"] += 1

    # 再处理图注：a) Caption 样式段落（pandoc 图片题注）；b) 紧贴图片的 "图：xxx" 段落
    for i, el in enumerate(children):
        p = as_para(el)
        if p is None or _has_seq(p):
            continue
        text = p.text.strip()
        if text.startswith("表"):
            continue
        prev = children[i - 1] if i > 0 else None
        style_name = ""
        try:
            style_name = p.style.name if p.style is not None else ""
        except Exception:
            pass
        is_caption_style = "Caption" in style_name and bool(text)
        is_fig_marker = bool(re.match(r"^图[：:]", text)) and has_drawing(prev)
        if is_caption_style or is_fig_marker:
            label = re.sub(r"^图[：:]\s*", "", text)
            _rewrite_caption(p, "图", label)
            if has_drawing(prev):  # 图片与其下方图注不分离
                prev_para = as_para(prev)
                if prev_para is not None:
                    prev_para.paragraph_format.keep_with_next = True
            stats["fig_captions"] += 1

    # 兜底：所有含图片的段落与后续内容保持同页（图片通常紧跟图注）
    for el in children:
        if el.tag == qn("w:p") and has_drawing(el):
            p = as_para(el)
            if p is not None:
                p.paragraph_format.keep_with_next = True
    return stats


EMU_PER_TWIP = 635  # 1 twip = 635 EMU


def _tblpr_insert_tblw(tblpr, tblw) -> None:
    """w:tblW 在 tblPr 里有固定次序（须在 jc/tblLayout/tblLook 之前）。"""
    later_tags = ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
                  "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    for tag in later_tags:
        anchor = tblpr.find(qn(tag))
        if anchor is not None:
            anchor.addprevious(tblw)
            return
    tblpr.append(tblw)


def polish_tables(doc: Document) -> dict:
    """表格字号统一为五号(10.5pt)，并按内容实测宽度定列宽、居中。

    pandoc 按 Markdown 源码字符数写死列宽（tblGrid/tcW），中文表头会被压成
    两行且与垂直居中的相邻列形成阶梯错位。这里：
      1) 表内文字设为 10.5pt（学术规范，也显著降低列宽压力）；
      2) 按 10.5pt 字符渲染宽度估算每列所需宽度（CJK/全角≈210 twips、
         字母数字≈118、其他≈105）+ 单元格边距，再留 12% 余量；
      3) 总宽不超版心则直接采用（单行零折行），超了等比收缩。
    显式写死宽度（Word 的 autofit 在 COM 导出时不可靠）。
    """
    stats = {"tables_polished": 0}
    content_twips = None
    if doc.sections:
        sec = doc.sections[0]
        content_twips = int((sec.page_width - sec.left_margin - sec.right_margin) / EMU_PER_TWIP)

    def est_twips(text: str) -> int:
        w = 0
        for ch in text:
            if ord(ch) > 0x2E7F:      # CJK / 全角标点 @10.5pt
                w += 210
            elif ch == " ":
                w += 62
            elif ch.isalnum():
                w += 118
            else:
                w += 105
        return w

    CELL_MARGINS = 260   # 单元格左右边距合计（twips）
    SAFETY = 1.12        # 列宽余量，避免 Word 两端对齐时临界折行
    MIN_COL = 900

    for table in doc.tables:
        try:
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            n_cols = len(table.columns)
            if n_cols == 0 or content_twips is None:
                continue
            # 表内文字统一五号；清零首行缩进（Compact 样式会继承
            # Body Text 的 480tw 首行缩进，导致单元格临界折行）
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.first_line_indent = Pt(0)
                        for run in p.runs:
                            run.font.size = Pt(10.5)

            required = [0] * n_cols
            for row in table.rows:
                cells = row.cells
                if len(cells) != n_cols:
                    continue
                for i, cell in enumerate(cells):
                    tw = min(int((est_twips(cell.text or "") + CELL_MARGINS) * SAFETY),
                             int(content_twips * 0.45))
                    required[i] = max(required[i], tw)
            required = [max(w, MIN_COL) for w in required]
            total = sum(required)
            if total <= content_twips:
                widths = required            # 每列单行放下，零折行
            else:
                widths = [max(int(w * content_twips / total), MIN_COL) for w in required]

            tbl = table._tbl
            grid = tbl.find(qn("w:tblGrid"))
            if grid is not None:
                for col_el, wd in zip(grid.findall(qn("w:gridCol")), widths):
                    col_el.set(qn("w:w"), str(wd))
            tblw = tbl.tblPr.find(qn("w:tblW"))
            if tblw is None:
                tblw = OxmlElement("w:tblW")
                _tblpr_insert_tblw(tbl.tblPr, tblw)
            tblw.set(qn("w:w"), str(sum(widths)))
            tblw.set(qn("w:type"), "dxa")
            for row in table.rows:
                cells = row.cells
                if len(cells) != n_cols:
                    continue
                for i, cell in enumerate(cells):
                    tcpr = cell._tc.get_or_add_tcPr()
                    tcw = tcpr.find(qn("w:tcW"))
                    if tcw is None:
                        tcw = OxmlElement("w:tcW")
                        tcpr.insert(0, tcw)  # tcW 在 tcPr 序列的最前部
                    tcw.set(qn("w:w"), str(widths[i]))
                    tcw.set(qn("w:type"), "dxa")
            table.autofit = False  # 固定布局，按上面写死的宽度渲染
            stats["tables_polished"] += 1
        except Exception:
            continue
    return stats


# ---------------------------------------------------------------- 主流程
def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("docx", help="要处理的 docx 路径（就地修改）")
    parser.add_argument("--heading-font", default="黑体")
    parser.add_argument("--body-font", default="宋体")
    parser.add_argument("--latin-font", default="Times New Roman")
    parser.add_argument("--number-levels", type=int, default=2, choices=(1, 2, 3))
    parser.add_argument("--toc-levels", default="1-3")
    parser.add_argument("--no-fonts", action="store_true")
    parser.add_argument("--no-number", action="store_true")
    parser.add_argument("--no-toc", action="store_true")
    parser.add_argument("--no-captions", action="store_true")
    ns = parser.parse_args()

    doc = Document(ns.docx)
    result: dict = {"file": ns.docx, "applied": [], "ok": True}
    try:
        # 注意顺序：目录/题注/编号会新建 run，字体处理必须最后执行才能覆盖它们
        if not ns.no_toc:
            result["toc"] = insert_toc(doc, ns.toc_levels)
            result["applied"].append("toc")
        if not ns.no_captions:
            result["captions"] = number_captions(doc)
            result["applied"].append("captions")
        result["tables"] = polish_tables(doc)
        result["applied"].append("tables_autofit_center")
        if not ns.no_number:
            result["numbering"] = number_headings(doc, ns.number_levels)
            result["applied"].append("heading_numbering")
        if not ns.no_fonts:
            stats = apply_fonts(doc, ns.heading_font, ns.body_font, ns.latin_font)
            result["applied"].append("fonts")
            result["fonts"] = stats
        doc.save(ns.docx)
    except Exception as exc:  # 让上层看到真实失败原因
        result["ok"] = False
        result["error"] = f"{type(exc).__name__}: {exc}"
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
