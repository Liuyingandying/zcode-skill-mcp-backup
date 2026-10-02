#!/usr/bin/env python3
"""交付前自动检查 v2.0：docx / pdf / pptx。生成后必须跑，结果如实报告给用户。

v2.0 新增（中国高校学术格式规范，配置来自 style_config.py）:
  页面设置（A4 + 页边距 2.5/2.5/3.0/2.5 cm）/ 页眉报告头 / 页脚页码域 /
  标题样式（黑体 + 三号/四号/小四 + 加粗）/ 正文样式（宋体小四 + 1.5 倍行距 +
  首行缩进 2 字符 + 西文 Times New Roman）/ 参考文献排版（有该章节时检查）

v1 检查项全部保留:
  docx : 能否打开 / 中文文本量 / eastAsia 字体覆盖率 / 目录域 / 标题编号 / 图片与"图 N"题注匹配 /
         表格与"表 N"题注匹配 / 是否含内容控件(SDT，WPS 兼容风险)
  pdf  : 能否打开 / 页数 / 首页中文可提取 / 内嵌字体含 CJK 字体
  pptx : 能否打开 / 页数 / 图片数 / 逐页字数(大段文字警告) / eastAsia 字体覆盖率

用法:
    python verify_output.py 文件1.docx 文件2.pdf [--style academic_paper]

输出: 每个文件一行 JSON {"file","type","ok","checks":[{name,pass,detail}]}；
全部通过退出码 0，任一失败 1。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from style_config import STYLE_NAMES, get_style  # noqa: E402

CJK_RE = re.compile(r"[\u4e00-\u9fff]")
CJK_FONT_RE = re.compile(
    r"(SimSun|SimHei|SimXi|SimFang|KaiTi|FangSong|YaHei|YaHeiUI|宋体|黑体|楷体|仿宋|雅黑"
    r"|Song|Hei|Kai|Fang|Ming|CJK|GBK|NotoSansCJK|NotoSerifCJK|SourceHan)", re.I)

EMU_PER_CM = 360000
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _check(name: str, passed: bool, detail: str) -> dict:
    return {"name": name, "pass": bool(passed), "detail": detail}


def _style_eastAsia(style) -> str | None:
    try:
        rpr = style.element.find(f"{W_NS}rPr")
        if rpr is None:
            return None
        rf = rpr.find(f"{W_NS}rFonts")
        return rf.get(f"{W_NS}eastAsia") if rf is not None else None
    except Exception:
        return None


# ---------------------------------------------------------------- v2.0 学术格式检查
def check_page_setup(doc: Document, cfg: dict) -> dict:
    page = cfg["page"]
    if not doc.sections:
        return _check("page_setup", False, "文档无 section")
    sec = doc.sections[0]
    tol = 0.06  # cm，EMU 取整容差
    got = {
        "width": sec.page_width / EMU_PER_CM,
        "height": sec.page_height / EMU_PER_CM,
        "top": sec.top_margin / EMU_PER_CM,
        "bottom": sec.bottom_margin / EMU_PER_CM,
        "left": sec.left_margin / EMU_PER_CM,
        "right": sec.right_margin / EMU_PER_CM,
    }
    mismatches = []
    for key, expect in (("width", page["width_cm"]), ("height", page["height_cm"]),
                        ("top", page["margin_top_cm"]), ("bottom", page["margin_bottom_cm"]),
                        ("left", page["margin_left_cm"]), ("right", page["margin_right_cm"])):
        if abs(got[key] - expect) > tol:
            mismatches.append(f"{key}={got[key]:.2f}(期望{expect})")
    ok = not mismatches
    return _check("page_setup", ok,
                  f"A4 {got['width']:.1f}x{got['height']:.1f}cm，边距 上{got['top']:.1f}/下{got['bottom']:.1f}"
                  f"/左{got['left']:.1f}/右{got['right']:.1f}cm" if ok else "；".join(mismatches))


def check_header_footer(doc: Document, cfg: dict) -> dict:
    if not doc.sections:
        return _check("header_footer", False, "文档无 section")
    sec = doc.sections[0]
    header_text = "\n".join(p.text for p in sec.header.paragraphs).strip()
    header_ok = cfg["header_text"] in header_text
    footer_xml = "".join(p._p.xml for p in sec.footer.paragraphs) if sec.footer is not None else ""
    footer_ok = "PAGE" in footer_xml and "instrText" in footer_xml
    ok = header_ok and footer_ok
    detail = f"页眉“{header_text[:20]}”" + ("✓" if header_ok else f"（期望含“{cfg['header_text']}”）") + \
             "；页脚页码域" + ("✓" if footer_ok else "缺失")
    return _check("header_footer", ok, detail)


def check_heading_styles(doc: Document, cfg: dict) -> dict:
    problems = []
    for level in (1, 2, 3):
        try:
            style = doc.styles[f"Heading {level}"]
        except KeyError:
            problems.append(f"Heading {level} 样式缺失")
            continue
        hcfg = cfg[f"heading{level}"]
        ea = _style_eastAsia(style)
        if ea != cfg["fonts"]["heading_eastasia"]:
            problems.append(f"H{level} 字体={ea}")
        size = style.font.size.pt if style.font.size is not None else None
        if size is None or abs(size - hcfg["size_pt"]) > 0.26:
            problems.append(f"H{level} 字号={size}")
        if not style.font.bold:
            problems.append(f"H{level} 未加粗")
    ok = not problems
    h1 = cfg["heading1"]["size_pt"]
    return _check("heading_styles", ok,
                  f"一至三级标题 {cfg['fonts']['heading_eastasia']}"
                  f"{h1:g}/{cfg['heading2']['size_pt']:g}/{cfg['heading3']['size_pt']:g}pt 加粗 ✓"
                  if ok else "；".join(problems))


def check_body_style(doc: Document, cfg: dict) -> dict:
    body = cfg["body"]
    problems = []

    styles_el = doc.styles.element
    dd_rpr = styles_el.find(f"{W_NS}docDefaults/{W_NS}rPrDefault/{W_NS}rPr")
    latin = None
    if dd_rpr is not None:
        rf = dd_rpr.find(f"{W_NS}rFonts")
        if rf is not None:
            latin = rf.get(f"{W_NS}ascii")
    if latin != cfg["fonts"]["latin"]:
        problems.append(f"文档默认西文={latin}")

    try:
        style = doc.styles["Body Text"]
    except KeyError:
        style = None
        problems.append("Body Text 样式缺失")
    if style is not None:
        ea = _style_eastAsia(style)
        if ea != cfg["fonts"]["body_eastasia"]:
            problems.append(f"正文中文字体={ea}")
        size = style.font.size.pt if style.font.size is not None else None
        if size is None or abs(size - body["size_pt"]) > 0.26:
            problems.append(f"正文字号={size}")
        ls = style.paragraph_format.line_spacing
        if ls is None or abs(ls - body["line_spacing"]) > 0.01:
            problems.append(f"行距={ls}")
        style_xml = style.element.xml
        fli = style.paragraph_format.first_line_indent
        indent_ok = (fli is not None and abs(fli.pt - body["size_pt"] * body["first_line_indent_chars"]) < 1.5) \
            or "firstLineChars" in style_xml
        if not indent_ok:
            problems.append("首行缩进缺失")

    ok = not problems
    return _check("body_style", ok,
                  f"正文 {cfg['fonts']['body_eastasia']}{body['size_pt']:g}pt/"
                  f"{cfg['fonts']['latin']}，1.5倍行距，首行缩进2字符 ✓" if ok else "；".join(problems))


def check_references(doc: Document, cfg: dict) -> dict:
    paras = doc.paragraphs
    start = None
    for i, p in enumerate(paras):
        try:
            name = p.style.name if p.style is not None else ""
        except Exception:
            continue
        if name.startswith("Heading") and "参考文献" in p.text:
            start = i + 1
            break
    if start is None:
        return _check("references_format", True, "无参考文献章节（跳过）")

    refs = []
    for p in paras[start:]:
        try:
            name = p.style.name if p.style is not None else ""
        except Exception:
            break
        if name.startswith("Heading") or name == "Title":
            break
        if p.text.strip():
            refs.append(p)
    if not refs:
        return _check("references_format", True, "参考文献章节为空（跳过）")

    problems = []
    numbered = sum(1 for p in refs if re.match(r"^\s*\[\d+\]", p.text))
    if numbered < len(refs):
        problems.append(f"仅 {numbered}/{len(refs)} 条带[n]编号")
    ok_size = ok_font = 0
    for p in refs:
        for run in p.runs:
            if not run.text:
                continue
            size = run.font.size.pt if run.font.size is not None else None
            rpr = run._r.find(f"{W_NS}rPr")
            ea = None
            if rpr is not None:
                rf = rpr.find(f"{W_NS}rFonts")
                ea = rf.get(f"{W_NS}eastAsia") if rf is not None else None
            if size is not None and abs(size - cfg["reference"]["size_pt"]) < 0.26:
                ok_size += 1
            if ea == cfg["fonts"]["body_eastasia"]:
                ok_font += 1
            break  # 每条只查首个非空 run
    if ok_size < len(refs):
        problems.append(f"仅 {ok_size}/{len(refs)} 条为五号")
    if ok_font < len(refs):
        problems.append(f"仅 {ok_font}/{len(refs)} 条显式宋体")
    ok = not problems
    return _check("references_format", ok,
                  f"{len(refs)} 条文献 [n]编号+五号宋体 ✓" if ok else "；".join(problems))


# ---------------------------------------------------------------- docx
def check_docx(path: str, strict_captions: bool = True, style_cfg: dict | None = None) -> dict:
    checks: list[dict] = []
    try:
        from docx import Document as _Doc
        doc = _Doc(path)
    except Exception as exc:
        return {"file": path, "type": "docx", "ok": False,
                "checks": [_check("open_ok", False, f"{type(exc).__name__}: {exc}")]}

    cfg = style_cfg or get_style("academic_paper")

    checks.append(_check("open_ok", True, "python-docx 解析成功"))

    # —— v2.0 学术格式 ——
    checks.append(check_page_setup(doc, cfg))
    checks.append(check_header_footer(doc, cfg))
    checks.append(check_heading_styles(doc, cfg))
    checks.append(check_body_style(doc, cfg))
    checks.append(check_references(doc, cfg))

    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.append(cell.text)
    cjk = sum(len(CJK_RE.findall(t)) for t in texts)
    checks.append(_check("cjk_text", cjk > 0, f"中文字符 {cjk} 个"))

    # eastAsia 字体覆盖率（含中文的 run 中已显式设置 eastAsia 的比例）
    total = fixed = 0
    doc_xml = doc.element.body.xml
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text and CJK_RE.search(run.text):
                total += 1
                rpr = run._r.find(f"{W_NS}rPr")
                if rpr is not None:
                    rf = rpr.find(f"{W_NS}rFonts")
                    if rf is not None and rf.get(f"{W_NS}eastAsia"):
                        fixed += 1
    if total == 0:
        checks.append(_check("eastasia_fonts", "docDefaults" in (doc.styles.element.xml or ""),
                             "正文无含中文 run；依赖样式/文档默认字体"))
    else:
        pct = round(100 * fixed / total)
        checks.append(_check("eastasia_fonts", pct >= 80, f"含中文 run 的 eastAsia 覆盖率 {pct}% ({fixed}/{total})"))

    toc_ok = "TOC" in doc_xml and "instrText" in doc_xml
    checks.append(_check("toc_field", toc_ok, "存在" if toc_ok else "无目录域（如为无目录的短文档可忽略）"))

    numbered, unnumbered = [], []
    for p in doc.paragraphs:
        try:
            name = p.style.name if p.style is not None else ""
        except Exception:
            continue
        if name in ("Heading 1", "Heading 2") and p.text.strip():
            (numbered if re.match(r"^\s*\d", p.text) else unnumbered).append(p.text.strip()[:20])
    if numbered or not unnumbered:
        checks.append(_check("heading_numbering", True,
                             f"已编号 {len(numbered)} 个一二级标题" + (f"；未编号 {len(unnumbered)} 个" if unnumbered else "")))
    else:
        checks.append(_check("heading_numbering", False, f"一二级标题均未编号，如：{unnumbered[:3]}"))

    images = len(doc.inline_shapes)
    fig_caps = doc_xml.count("SEQ 图")
    tab_caps = doc_xml.count("SEQ 表")
    tables = len(doc.tables)
    if images == 0:
        checks.append(_check("figure_captions", True, "无图片"))
    else:
        ok = fig_caps == images
        checks.append(_check("figure_captions", ok if strict_captions else True,
                             f"图片 {images} 张，图题注 {fig_caps} 个" + ("" if ok else "（数量不匹配）")))
    if tables == 0:
        checks.append(_check("table_captions", True, "无表格"))
    else:
        ok = tab_caps == tables
        checks.append(_check("table_captions", ok if strict_captions else True,
                             f"表格 {tables} 个，表题注 {tab_caps} 个" + ("" if ok else "（数量不匹配）")))

    sdt = doc_xml.count("<w:sdt>")
    checks.append(_check("wps_compat", sdt == 0, "无内容控件(SDT)" if sdt == 0 else f"发现 {sdt} 个内容控件，WPS 兼容风险"))

    return {"file": path, "type": "docx", "ok": all(c["pass"] for c in checks), "checks": checks}


# ---------------------------------------------------------------- pdf
def check_pdf(path: str) -> dict:
    checks: list[dict] = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
    except Exception as exc:
        return {"file": path, "type": "pdf", "ok": False,
                "checks": [_check("open_ok", False, f"{type(exc).__name__}: {exc}")]}

    checks.append(_check("open_ok", True, f"共 {len(reader.pages)} 页"))

    fonts: set[str] = set()
    for page in reader.pages:
        try:
            res = page.get("/Resources", {})
            for f in (res.get("/Font", {}) or {}).values():
                obj = f.get_object()
                name = str(obj.get("/BaseFont", ""))
                if name:
                    fonts.add(name.split("+")[-1])
        except Exception:
            continue
    cjk_fonts = [f for f in fonts if CJK_FONT_RE.search(f)]
    checks.append(_check("cjk_fonts", bool(cjk_fonts),
                         f"内嵌 CJK 字体: {sorted(set(cjk_fonts))[:4]}" if cjk_fonts else f"未检出 CJK 字体（全部字体: {sorted(fonts)[:6]}）"))

    sample = ""
    for page in reader.pages[:3]:
        try:
            sample = page.extract_text() or ""
        except Exception:
            sample = ""
        if CJK_RE.search(sample):
            break
    checks.append(_check("cjk_text", bool(CJK_RE.search(sample)),
                         "中文可正常提取" if CJK_RE.search(sample) else "前 3 页未提取到中文，可能字体/编码异常"))

    return {"file": path, "type": "pdf", "ok": all(c["pass"] for c in checks), "checks": checks}


# ---------------------------------------------------------------- pptx
def check_pptx(path: str, max_chars: int) -> dict:
    checks: list[dict] = []
    try:
        from pptx import Presentation
        prs = Presentation(path)
    except Exception as exc:
        return {"file": path, "type": "pptx", "ok": False,
                "checks": [_check("open_ok", False, f"{type(exc).__name__}: {exc}")]}

    checks.append(_check("open_ok", True, f"共 {len(prs.slides)} 页"))

    total_chars = 0
    heavy: list[str] = []
    images = 0
    runs_total = runs_ea = 0

    def walk(shapes, idx: int):
        nonlocal total_chars, images, runs_total, runs_ea
        for shape in shapes:
            if shape.shape_type == 6:
                walk(shape.shapes, idx)
                continue
            if shape.shape_type == 13:  # PICTURE
                images += 1
            if not getattr(shape, "has_text_frame", False):
                continue
            page_text = ""
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    page_text += run.text
                    if run.text and CJK_RE.search(run.text):
                        runs_total += 1
                        rpr = run._r.find(
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                        if rpr is not None and rpr.find(
                                "{http://schemas.openxmlformats.org/drawingml/2006/main}ea") is not None:
                            runs_ea += 1
            page_chars = len(CJK_RE.findall(page_text)) + len(re.findall(r"[A-Za-z]+", page_text))
            if page_chars > max_chars:
                heavy.append(f"第{idx}页约{page_chars}字")
            total_chars += page_chars

    for i, slide in enumerate(prs.slides, start=1):
        walk(slide.shapes, i)
    checks.append(_check("slide_density", not heavy,
                         "每页文字量正常" if not heavy else f"大段文字警告: {'; '.join(heavy)}（应压缩为要点，完整内容移到演讲备注）"))
    checks.append(_check("images", True, f"图片/图表对象 {images} 个，全篇文字量约 {total_chars} 字"))
    pct = round(100 * runs_ea / runs_total) if runs_total else 100
    checks.append(_check("eastasia_fonts", pct >= 80,
                         f"含中文 run 的 eastAsia 覆盖率 {pct}%" + ("" if pct >= 80 else "（先运行 fix_pptx_fonts.py）")))
    return {"file": path, "type": "pptx", "ok": all(c["pass"] for c in checks), "checks": checks}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("files", nargs="+")
    parser.add_argument("--style", default="academic_paper", choices=list(STYLE_NAMES))
    parser.add_argument("--max-slide-chars", type=int, default=100)
    ns = parser.parse_args()

    style_cfg = get_style(ns.style)
    all_ok = True
    for path in ns.files:
        lower = path.lower()
        if lower.endswith(".docx"):
            report = check_docx(path, style_cfg=style_cfg)
        elif lower.endswith(".pdf"):
            report = check_pdf(path)
        elif lower.endswith(".pptx"):
            report = check_pptx(path, ns.max_slide_chars)
        else:
            report = {"file": path, "type": "?", "ok": False,
                      "checks": [_check("supported", False, "仅支持 docx/pdf/pptx")]}
        all_ok &= report["ok"]
        print(json.dumps(report, ensure_ascii=False))
    return 0 if all_ok else 1


if __name__ == "__main__":
    sys.exit(main())
